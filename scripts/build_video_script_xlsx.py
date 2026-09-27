"""把视频分镜脚本的旁白整理成 Excel。

数据来源：docs/VIDEO_SCRIPT_2330_3000.md
输出：docs/VIDEO_SCRIPT_2330_3000.xlsx

五列：镜头编号 / 镜头内容(简述) / 时长 / 英文旁白 / 中文旁白
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "VIDEO_SCRIPT_2330_3000.xlsx"

# (镜头编号, 镜头内容简述, 时长, 英文旁白, 中文旁白)
ROWS = [
    ("镜头 1", "工具清单里没有审批权（终端）", "18 秒",
     "Let us start with the safety boundary. These are all the tools the agent can call — notice there is no approve, no reject, no apply. The model can only read context and produce proposals; approval is never the model's authority.",
     "先看安全边界。这是 Agent 能调用的全部工具——注意，没有 approve、没有 reject、没有 apply。模型只能读取上下文和生成提议，批准从来不是模型的权限。"),
    ("镜头 2", "信息不足时先追问（浏览器：Customer）", "25 秒",
     "The resident only says there is water on the floor. That could be an air-conditioner leak or a pipe leak. The system does not guess — it asks which one, with examples the resident can follow. When key information is missing, it asks first instead of scheduling.",
     "住户只说地上有水。这句话既可能是空调漏水，也可能是水管漏水。系统不会替他猜——它会反问是哪一种，并给出可以照着回答的例子。信息不足时先追问，而不是先排班。"),
    ("镜头 3", "硬约束与排除原因（浏览器：Coordinator）", "25 秒",
     "On the coordinator side, every candidate technician is evaluated: who qualifies, who is excluded, and why — a skill mismatch, a missing certification, or a time conflict. A technician who violates a hard constraint is never recommended, and that decision is made by deterministic code, not the model's impression.",
     "协调员这边能看到每一个候选技术员的评估结果：谁合格、谁被排除、以及为什么——技能不符、缺少认证，还是时段冲突。不满足硬约束的技术员不会被推荐，这个判断由确定性代码做出，不是模型的印象。"),
    ("镜头 4", "推荐理由与取舍（浏览器：Coordinator）", "20 秒",
     "The recommendation rationale comes from stored decision facts: why this technician, how the workload rises, and the ranking that was applied. The explanation is verifiable, not the model talking.",
     "推荐理由来自存储的决策事实：为什么选这个人、工作量会从多少涨到多少、排序依据是什么。解释是可核对的，不是模型的说辞。"),
    ("镜头 5", "Agent 无权批准高影响变更（浏览器：Coordinator）", "20 秒",
     "When a technician reports leave or a delay, the system produces a recovery plan. Every change is labelled: whether the technician changed, whether the time changed, whether approval is required, and why. Right now the schedule has not been touched — this is only a proposal. High-impact changes must be approved by a human.",
     "技术员报告请假或延误后，系统会生成恢复方案。每条变更都标注是否换人、是否改时间、是否需要人工审批以及原因。此刻排班还没有被改动——这只是一个提议。高影响的操作必须由人批准。"),
    ("镜头 6", "可审计与草稿通知（浏览器：Coordinator）", "12 秒",
     "Tool calls, handoffs, and the approval process are fully auditable. Customer notifications are only drafts — they are never sent automatically.",
     "工具调用、handoff 与审批过程全程可审计。客户通知也只是草稿，不会自动发送。"),
    ("镜头 7", "产品测试（终端）", "15 秒",
     "First, product tests: one hundred and seventy-six tests all pass, covering intake understanding, scheduling, disruption recovery, governance, and multimodal intake.",
     "先看产品测试：一百七十六个用例全部通过，覆盖意图理解、调度、改期、治理与多模态。"),
    ("镜头 8", "离线基准 120/120（终端）", "30 秒",
     "Next, our separate system benchmark — one hundred and twenty offline cases. Request understanding, schedule feasibility, and minimum-change recovery all score full marks; unauthorized and pre-approval schedule mutations in the governance layer are all zero; and the safe-failure rate under adversarial cases is one hundred percent. These measure deterministic correctness and safety.",
     "接着是独立的系统基准，一百二十个离线用例。请求理解、调度可行性、改期最小改动全部满分；治理层的越权变更与审批前修改全部为零；对抗场景的安全失败率百分之百。这些衡量的是确定性正确性与安全。"),
    ("镜头 9", "Live 基准单独呈现（编辑器 / 终端）", "25 秒",
     "Finally, the live benchmark, run against the real Claude Sonnet 4.5 on forty-five cases: task completion is one hundred percent, orchestration and handoff accuracy is eighty percent, p50 latency is about thirty seconds and p95 about thirty-three. To be clear: the offline full marks are deterministic correctness, while live measures the real model's orchestration quality — two different dimensions, and we do not conflate them.",
     "最后是 Live 基准，用真实的 Claude Sonnet 四点五跑四十五个用例：任务完成率百分之百，编排与 handoff 准确率百分之八十，p50 延迟约三十秒，p95 约三十三秒。必须强调：离线的满分是确定性正确性，Live 是真实模型的编排质量，两者维度不同，我们不混为一谈。"),
    ("镜头 10", "证据到底证明了什么（承接画面）", "20 秒",
     "This evidence proves four things: no illegal schedule is created; affected jobs are correctly detected; the schedule is not mutated before approval; and when no feasible plan exists, the workflow stops safely and hands off to a human. One live edge case routed differently than expected, but every safety invariant still held — which is exactly why we put safety in deterministic code rather than in the model.",
     "这些证据证明四件事：不会产生非法排班；能正确识别受影响的工单；审批之前不修改排班；当没有可行方案时工作流安全停止并转人工。Live 里有一个边界用例的路由与预期不同，但所有安全不变式依然保持——这正是我们把安全放在确定性代码而不是模型里的原因。"),
    ("镜头 11", "最小启动数据（编辑器）", "35 秒",
     "This is all the data a pilot needs: technicians with their skills and certifications, working hours, existing bookings, service rules, and a small set of customer requests. There is nothing else to prepare.",
     "试点需要的数据就这些：技术员及其技能与认证、工作时间、已有预约、服务规则，再加一小批客户请求。没有别的前置条件。"),
    ("镜头 12", "几条命令完成初始化（终端）", "25 秒",
     "Once these tables are imported, one init command plus one validation command, and the system is ready to take requests. Mendigo is a lightweight standalone pilot: an SME does not have to replace its existing systems or finish a large-scale digital transformation before testing the core workflow.",
     "把这些表导入后，一条初始化命令加一条校验命令，系统就可以开始接请求。Mendigo 是一个轻量的独立试点：不要求 SME 替换现有系统，也不要求先完成大规模数字化改造，就能验证核心流程。"),
    ("镜头 13", "诚实的边界（幻灯片或承接终端画面）", "30 秒",
     "Let us be honest about the boundary: today this step is assisted by the deployment team, not self-service onboarding. But it needs no enterprise ERP — starting with one service category, a small team, and one coordinator is enough.",
     "需要说清楚现实边界：今天这一步由部署团队协助完成，还不是自助式上线。但它不需要企业级 ERP——从一个服务类别、一个小团队、一位协调员开始就够了。"),
    ("镜头 14", "已实现 vs 未来（幻灯片，左右两栏）", "45 秒",
     "What is implemented and demonstrable today: request intake and clarification, technician recommendation and coordinator approval, the technician schedule view, unavailable and delayed disruption recovery, plus a full audit trail and a quantitative benchmark. Future milestones include self-service data import and system connectors, travel-time buffers with live traffic, real notification sending, and individual user accounts. We only show what we have actually built.",
     "目前已实现并可演示的是：请求接入与澄清、技术员推荐与协调员审批、技术员排班视图、请假与延误两类改期恢复，以及完整的审计轨迹和量化基准。后续里程碑包括自助数据导入与系统对接、通勤时间缓冲与实时路况、真实通知发送，以及个人账号体系。我们只展示已经做到的部分。"),
    ("镜头 15", "Before / After 对比（幻灯片）", "25 秒",
     "Before: fragmented messages, manual checking, and chaos the moment something changes. After: a structured job, an explainable recommendation, humans in control of the key decisions, and fast recovery when plans change.",
     "改变之前：消息零散、人工核对、一旦出现变动就陷入混乱。改变之后：工单结构化、推荐可解释、关键决定由人掌控、变动后能快速恢复。"),
    ("镜头 16", "收束（幻灯片）", "20 秒",
     "Mendigo keeps every service promise — even when the day does not go as planned.",
     "Mendigo 让每一个服务承诺都被兑现——即使这一天并没有按计划进行。"),
]

HEADERS = ["镜头编号", "镜头内容（简述）", "时长", "英文旁白", "中文旁白"]

wb = Workbook()
ws = wb.active
ws.title = "视频旁白 2330-3000"

# 表头
header_fill = PatternFill("solid", fgColor="305496")
header_font = Font(bold=True, color="FFFFFF", size=11)
for col, title in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=col, value=title)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# 数据行
for r, row in enumerate(ROWS, start=2):
    for c, value in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=value)
        # 编号/时长居中，旁白与简述左对齐并自动换行
        if c in (1, 3):
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

# 列宽
widths = {"A": 10, "B": 34, "C": 8, "D": 70, "E": 55}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# 冻结表头，加筛选
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(ROWS)+1}"

wb.save(OUT)
print(f"已生成：{OUT}  （共 {len(ROWS)} 个镜头）")
